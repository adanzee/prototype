
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Optional
import hashlib
import json
import re

import pymupdf
from docx import Document


@dataclass(frozen=True)
class TextBlock:
    block_id: str
    text: str
    source_type: str
    page_number: Optional[int]
    section: Optional[str] = None


class DocumentParser:
    SUPPORTED_EXTENSIONS = {".txt", ".docx", ".pdf"}

    def parse(self, file_path: str | Path) -> list[TextBlock]:
        path = Path(file_path)

        if not path.is_file():
            raise FileNotFoundError(f"File not found: {path}")

        extension = path.suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {extension}")

        if extension == ".txt":
            raw_blocks = self._parse_txt(path)
        elif extension == ".docx":
            raw_blocks = self._parse_docx(path)
        else:
            raw_blocks = self._parse_pdf(path)

        blocks = []

        for index, (text, source_type, page_number) in enumerate(raw_blocks):
            text = text.strip()

            if not text:
                continue

            blocks.append(
                TextBlock(
                    block_id=self._make_block_id(
                        path.name,
                        source_type,
                        page_number,
                        index,
                        text,
                    ),
                    text=text,
                    source_type=source_type,
                    page_number=page_number,
                )
            )

        return blocks

    def _parse_txt(self, path: Path) -> list[tuple[str, str, Optional[int]]]:
        content = path.read_text(encoding="utf-8-sig")

        paragraphs = re.split(r"\n\s*\n", content)

        return [
            (paragraph.strip(), "txt", None)
            for paragraph in paragraphs
            if paragraph.strip()
        ]

    def _parse_docx(self, path: Path) -> list[tuple[str, str, Optional[int]]]:
        document = Document(path)
        blocks = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                blocks.append((text, "docx", None))

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    blocks.append((" | ".join(cells), "docx", None))

        return blocks

    def _parse_pdf(self, path: Path) -> list[tuple[str, str, Optional[int]]]:
        blocks = []

        with pymupdf.open(path) as document:
            for page_number, page in enumerate(document, start=1):
                for block in page.get_text("blocks"):
                    text = block[4].strip()

                    if text:
                        blocks.append((text, "pdf", page_number))

        return blocks

    @staticmethod
    def _make_block_id(
        filename: str,
        source_type: str,
        page_number: Optional[int],
        index: int,
        text: str,
    ) -> str:
        content = f"{filename}|{source_type}|{page_number}|{index}|{text}"
        digest = hashlib.sha256(content.encode("utf-8")).hexdigest()[:16]

        return f"block_{digest}"

    @staticmethod
    def to_json(blocks: list[TextBlock]) -> str:
        return json.dumps(
            [asdict(block) for block in blocks],
            ensure_ascii=False,
            indent=2,
        )
